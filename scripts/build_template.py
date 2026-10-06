"""生成 Word 模板 templates/design-report-v1.docx（ADR-0005，章节规格见 docs/10）。

模板生成后可直接在 Word 中调整版式；本脚本保证可重复生成。
docxtpl 表格循环约定：`{%tr for %}` 标签行与 `{%tr endfor %}` 行只做控制，
**两者之间的内容行**才会被重复；因此每个循环表为「表头 / 标签行 / 内容行 / endfor 行」四行。
"""
from __future__ import annotations

from pathlib import Path

from docx import Document

BASE = Path(__file__).resolve().parent.parent
OUT = BASE / "templates" / "design-report-v1.docx"


def _loop_table(doc: Document, headers: list[str], loop_expr: str, body_cells: list[str]) -> None:
    """表头 + 控制标签行 + 循环内容行 + endfor 行。"""
    table = doc.add_table(rows=4, cols=len(headers))
    table.style = "Table Grid"
    for index, title in enumerate(headers):
        table.rows[0].cells[index].text = title
    table.rows[1].cells[0].text = "{%tr for " + loop_expr + " %}"
    for index, cell_text in enumerate(body_cells):
        table.rows[2].cells[index].text = cell_text
    table.rows[3].cells[0].text = "{%tr endfor %}"


def build() -> Path:
    doc = Document()
    doc.add_heading("卫星整星详细设计方案", level=0)
    doc.add_paragraph("模板：design-report-v1 ｜ 由 AI 卫星总体设计助手生成")

    # 1 概述
    doc.add_heading("1 概述", level=1)
    doc.add_paragraph("{{ narrative.overview }}")

    # 2 总体构型
    doc.add_heading("2 卫星总体构型", level=1)
    doc.add_paragraph("构型类型：{{ config.type }}。{{ config.notes }}")
    table = doc.add_table(rows=4, cols=2)
    table.style = "Table Grid"
    table.rows[0].cells[0].text = "包络尺寸（mm）"
    table.rows[0].cells[1].text = "值"
    for index, axis in enumerate(("x", "y", "z"), start=1):
        table.rows[index].cells[0].text = axis.upper()
        table.rows[index].cells[1].text = "{{ envelope.%s }}" % axis

    # 3 分系统配置
    doc.add_heading("3 分系统配置", level=1)
    _loop_table(doc, ["分系统", "说明"], "s in subsystems", ["{{ s.name }}", "{{ s.description }}"])

    # 4~6 预算
    for number, title, loop_key, narrative_key in (("4", "质量预算", "mass_items", "mass"),
                                                   ("5", "功耗预算", "power_items", "power"),
                                                   ("6", "数据量预算", "data_items", "data")):
        doc.add_heading(f"{number} {title}", level=1)
        doc.add_paragraph("{{ narrative.%s }}" % narrative_key)
        _loop_table(
            doc,
            ["项目", "数值", "单位", "来源"],
            f"it in {loop_key}",
            ["{{ it.name }}", "{{ it.value }}", "{{ it.unit }}", "{{ it.source_ref }}"],
        )

    # 7 测控数传
    doc.add_heading("7 测控数传方案", level=1)
    doc.add_paragraph("链路预算（自由空间损耗 / Eb-N0 / 余量）：")
    _loop_table(
        doc,
        ["项目", "数值", "单位", "来源"],
        "it in link_items",
        ["{{ it.name }}", "{{ it.value }}", "{{ it.unit }}", "{{ it.source_ref }}"],
    )

    # 8 姿控
    doc.add_heading("8 姿态确定与控制方案", level=1)
    doc.add_paragraph("姿态控制方案采用三轴稳定 + 反作用轮 + 磁力矩器（详细指标待姿控模块接入后补全）。")

    # 9 轨道与覆盖
    doc.add_heading("9 轨道与覆盖分析", level=1)
    doc.add_paragraph("轨道类型 {{ mission.orbit_type }}，轨道高度 {{ mission.altitude_km }} km，"
                      "设计寿命 {{ mission.lifetime_years }} 年。")
    _loop_table(
        doc,
        ["项目", "数值", "单位", "来源"],
        "it in orbit_items",
        ["{{ it.name }}", "{{ it.value }}", "{{ it.unit }}", "{{ it.source_ref }}"],
    )
    doc.add_paragraph("模型说明：本版采用二体 + J2 长期项简化模型；覆盖与链路结果为任务级估算，"
                      "不适用于高精度定轨与链路设计，误差范围见《计算与校验引擎设计》。")

    # 10 选型
    doc.add_heading("10 关键器件与选型建议", level=1)
    doc.add_paragraph("{% if not candidates %}候选清单待知识库模块接入后生成。{% endif %}")

    # 11 风险
    doc.add_heading("11 风险分析", level=1)
    doc.add_paragraph("{{ narrative.risk }}")

    # 12 引用与假设
    doc.add_heading("12 引用与假设", level=1)
    doc.add_paragraph("引用来源：")
    _loop_table(doc, ["编号", "来源", "链接", "类型"], "c in citations",
                ["{{ c.id }}", "{{ c.title }}", "{{ c.url }}", "{{ c.type }}"])
    doc.add_paragraph("{% if not citations %}（暂无引用：知识库模块接入后自动填充）{% endif %}")

    doc.add_paragraph("假设台账：")
    _loop_table(doc, ["编号", "内容", "依据", "状态"], "a in assumptions",
                ["{{ a.id }}", "{{ a.content }}", "{{ a.basis }}", "{{ a.status }}"])

    doc.add_paragraph("本文件由系统自动生成；全部数值来自确定性计算或显式假设，出处见本附录。")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT))
    return OUT


if __name__ == "__main__":
    path = build()
    print(f"模板已生成：{path}")
