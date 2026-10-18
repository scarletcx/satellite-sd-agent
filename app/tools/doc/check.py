"""渲染自检（docs/10 §4）：占位符残留 + 数字反查。

数字反查：文档中的每个数值必须能在「允许集合」（IR 参数 / 任务约束 / 假设台账 /
系统方案）中找到对应值（含舍入容差）；章节编号与年份不计。
"""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document

PLACEHOLDER_RE = re.compile(r"\{\{|\}\}")
NUMBER_RE = re.compile(r"\d+(?:\.\d+)?")
# 标识符类 token（来源编号 / 步骤号 / 物理符号）不参与数字反查：
# 如 calc-20261005-0013、asm-0001、t-20261005-001、S1、J2、AM0
IDENTIFIER_RE = re.compile(r"[A-Za-z]+(?:-\d+)+|[A-Za-z]+\d+")


def numbers_in_texts(texts: list[str]) -> list[float]:
    values: list[float] = []
    for text in texts:
        values.extend(float(token) for token in NUMBER_RE.findall(text or ""))
    return values


def _iter_texts(document: Document):
    for paragraph in document.paragraphs:
        yield paragraph.text
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                yield cell.text


def _excluded(value: float) -> bool:
    """章节编号（1~12）与年份（2000~2100）不参与数字反查。"""
    if not float(value).is_integer():
        return False
    return 1 <= value <= 12 or 2000 <= value <= 2100


def extract_texts(docx_path: Path) -> list[str]:
    """读取 docx 全文（段落 + 表格单元格），供数字反查与引用检查共用。"""
    document = Document(str(docx_path))
    return [text for text in _iter_texts(document)]


CITATION_REF_RE = re.compile(r"\bS\d+\b")
# 「S4 草案」这类步骤标签不是引用（假设台账、说明文字里会出现）
_STEP_LABEL_TAILS = ("草案", "步骤", "阶段", "设计输入")


def check_citations(texts: list[str], citations: list[dict]) -> dict:
    """引用检查：文中引用的编号必须真实存在；引用条目必须有来源标题。

    只把真正的引用当作引用：形如 `[S1]`、`（S1）` 或独立出现的 S 编号；
    「（S4 草案）」这类步骤标签会被排除。
    """
    known = {str(item.get("id")) for item in citations if item.get("id")}
    referenced: set[str] = set()
    for text in texts:
        content = text or ""
        for match in CITATION_REF_RE.finditer(content):
            tail = content[match.end(): match.end() + 4].lstrip()
            if tail.startswith(_STEP_LABEL_TAILS):
                continue
            referenced.add(match.group(0))

    issues: list[dict] = [
        {"type": "citation_unknown", "detail": ref}
        for ref in sorted(referenced) if ref not in known
    ]
    issues += [
        {"type": "citation_missing_source", "detail": str(item.get("id"))}
        for item in citations if not str(item.get("title") or "").strip()
    ]
    return {"referenced": sorted(referenced), "known": sorted(known), "issues": issues}


def check_document(docx_path: Path, allowed_numbers: list[float]) -> dict:
    document = Document(str(docx_path))
    text = "\n".join(_iter_texts(document))

    placeholders_left = len(PLACEHOLDER_RE.findall(text))
    scan_text = IDENTIFIER_RE.sub(" ", text)  # 剥离标识符后再做数字反查

    checked = traceable = 0
    issues: list[dict] = []
    for match in NUMBER_RE.finditer(scan_text):
        value = float(match.group(0))
        if _excluded(value):
            continue
        checked += 1
        if any(abs(value - allowed) <= max(0.05, abs(allowed) * 0.001) for allowed in allowed_numbers):
            traceable += 1
        else:
            issues.append({"type": "number_untraceable", "detail": match.group(0)})

    if placeholders_left:
        issues.append({"type": "placeholder_left", "detail": f"{placeholders_left} 处未渲染占位符"})

    return {
        "placeholders_left": placeholders_left,
        "numbers_checked": checked,
        "numbers_traceable": traceable,
        "issues": issues,
        "status": "pass" if not issues else "warn",
    }
