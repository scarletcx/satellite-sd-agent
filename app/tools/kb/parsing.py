"""文档解析（docs/04 §3）：PDF（PyMuPDF）/ DOCX / TXT / MD / CSV → [(page, text)]。"""
from __future__ import annotations

from pathlib import Path

SUPPORTED_SUFFIXES = {".txt", ".md", ".csv", ".pdf", ".docx"}


def parse_file(path: Path) -> list[tuple[int | None, str]]:
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise ValueError(f"不支持的文件类型：{suffix}（支持 {sorted(SUPPORTED_SUFFIXES)}）")

    if suffix in {".txt", ".md", ".csv"}:
        return [(None, path.read_text(encoding="utf-8", errors="ignore"))]

    if suffix == ".pdf":
        import pymupdf

        with pymupdf.open(str(path)) as pdf:
            return [(index + 1, page.get_text()) for index, page in enumerate(pdf)]

    # .docx（表格按行拼接为文本）
    from docx import Document

    document = Document(str(path))
    parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return [(None, "\n".join(parts))]
