"""端到端流水线测试（桩模式，docs/11 §3 E2E-01 / ADV-04 门禁路径）。

流程含两道门禁：S4 后的设计输入确认（draft/未确认假设）与 S6 后的超限处置。
"""
import asyncio
import json
from datetime import datetime, timezone

import pytest
from docx import Document

from app.db import init_db, new_session
from app.models import (
    Assumption,
    DesignParameter,
    Mission,
    ReviewGate,
    SelectionCandidate,
    ValidationResult,
)
from app.orchestrator import engine
from app.orchestrator.ids import next_mission_id
from app.tracing import read_trace, task_dir
from tests.seed_kb import seed as seed_kb


def _create_mission(constraints: dict) -> str:
    with new_session() as session:
        task_id = next_mission_id(session)
        session.add(Mission(
            id=task_id,
            goal="设计一颗 500km SSO 光学遥感小卫星",
            constraints=constraints,
            status="queued",
        ))
        session.commit()
    return task_id


def _confirm_inputs(task_id: str) -> None:
    """模拟工程师在输入门禁中确认全部 S4 草案与答案假设。"""
    with new_session() as session:
        now = datetime.now(timezone.utc)
        for row in session.query(Assumption).filter_by(mission_id=task_id, status="open").all():
            row.status = "confirmed"
            row.confirmed_by = "tester"
            row.confirmed_at = now
        drafts = session.query(DesignParameter).filter_by(mission_id=task_id, status="draft").all()
        base = session.query(Assumption).filter_by(mission_id=task_id).count()
        for offset, row in enumerate(drafts, start=1):
            assumption_id = f"asm-{base + offset:04d}"
            session.add(Assumption(
                id=assumption_id, mission_id=task_id,
                content=f"设计输入「{row.name}」经 tester 确认",
                basis="测试确认", related_params=[row.id], status="confirmed",
                registered_by="review:tester", confirmed_by="tester", confirmed_at=now,
            ))
            row.status = "assumed"
            row.source_type = "assumption"
            row.source_ref = assumption_id
        session.commit()


def _close_gates(task_id: str) -> None:
    with new_session() as session:
        gates = (session.query(ReviewGate)
                 .filter(ReviewGate.mission_id == task_id, ReviewGate.action.is_(None)).all())
        for gate in gates:
            gate.action = "accept"
            gate.reviewer = "tester"
            gate.ts = datetime.now(timezone.utc)
        session.commit()


def test_pipeline_end_to_end():
    init_db()
    seed_kb()
    task_id = _create_mission({"orbit_type": "SSO", "altitude_km": 500,
                               "lifetime_years": 3, "mass_limit_kg": 80})
    asyncio.run(engine.run_mission(task_id))

    # 第一道门禁：S4 草案输入与假设必须人工确认
    with new_session() as session:
        mission = session.get(Mission, task_id)
        assert mission.status == "awaiting_gate", mission.last_error
        assert mission.current_step == "S4"
        drafts = session.query(DesignParameter).filter_by(mission_id=task_id, status="draft").all()
        assert len(drafts) == 6, "6 组设计输入应全部以 draft 状态落库"
        assert all(row.payload for row in drafts), "草案应携带完整分组 payload"
        assert session.query(Assumption).filter_by(mission_id=task_id, status="open").count() >= 1

    _confirm_inputs(task_id)
    _close_gates(task_id)
    asyncio.run(engine.continue_mission(task_id))

    candidate_names: list[str] = []
    with new_session() as session:
        mission = session.get(Mission, task_id)
        assert mission.status == "succeeded", mission.last_error

        params = {p.id: p for p in session.query(DesignParameter)
                  .filter_by(mission_id=task_id).all()}
        assert len(params) >= 28, "6 组输入 + 22 个计算参数"
        assert params["input.mass"].status == "assumed"
        assert params["input.mass"].source_type == "assumption"
        assert params["mass.total"].value > 0
        assert params["mass.total_with_margin"].value > params["mass.total"].value
        assert params["mass.total_with_margin"].status == "verified"
        assert params["power.battery_capacity_wh"].value > 0
        # C4/C5/C6 已接入：轨道、太阳翼、链路
        assert params["orbit.period_min"].value == pytest.approx(94.6, abs=1.5)
        assert params["orbit.orbits_per_day"].value > 15
        assert params["power.sa.area"].value > 0
        assert params["link.margin_db"].value > 3
        # C7 姿控粗算（依赖质量/面积/周期）
        assert params["adcs.h_required_nms"].value > 0
        assert params["adcs.mtq_dipole_am2"].value > 0

        results = session.query(ValidationResult).filter_by(mission_id=task_id).all()
        assert results and all(r.status != "block" for r in results)

        candidates = session.query(SelectionCandidate).filter_by(mission_id=task_id).all()
        assert len(candidates) >= 4, "选型库应产出多类候选"
        assert {"battery", "solar_cell", "transponder"} <= {c.category for c in candidates}
        assert all(c.reason for c in candidates), "候选应携带选择理由"
        candidate_names = [c.name for c in candidates]

    directory = task_dir(task_id)
    report = json.loads((directory / "check-report.json").read_text(encoding="utf-8"))
    assert report["placeholders_left"] == 0
    assert report["status"] == "pass"
    assert report["numbers_checked"] > 0
    assert (directory / "document-v1.docx").exists()

    # 引用链路：检索 → citations.json → 文档附录
    citations = json.loads((directory / "citations.json").read_text(encoding="utf-8"))
    assert citations, "S2 应检索到引用"
    assert report["citations"] == len(citations)
    document = Document(str(directory / "document-v1.docx"))
    texts = [p.text for p in document.paragraphs]
    texts += [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    assert citations[0]["title"] in "\n".join(texts), "附录应包含引用来源"
    assert candidate_names and candidate_names[0] in "\n".join(texts), "第 10 章应包含选型候选"

    events = read_trace(task_id)
    assert len(events) >= 9
    assert any(event.get("type") == "gate" for event in events), "trace 应包含门禁事件"


def test_pipeline_gate_on_mass_limit():
    init_db()
    seed_kb()
    task_id = _create_mission({"orbit_type": "SSO", "altitude_km": 500,
                               "mass_limit_kg": 70})
    asyncio.run(engine.run_mission(task_id))          # 停在输入门禁
    _confirm_inputs(task_id)
    _close_gates(task_id)
    asyncio.run(engine.continue_mission(task_id))     # S5/S6 → 质量超限门禁

    with new_session() as session:
        mission = session.get(Mission, task_id)
        assert mission.status == "awaiting_gate", mission.last_error
        gate = (session.query(ReviewGate)
                .filter(ReviewGate.mission_id == task_id, ReviewGate.action.is_(None))
                .first())
        assert gate is not None and gate.type == "discrepancy_resolve"
        param = session.get(DesignParameter, ("mass.total_with_margin", task_id))
        assert param.status == "needs_review"
        # 人工放行（模拟 A8 accept）
        gate.action = "accept"
        gate.reviewer = "tester"
        param.status = "verified"
        session.commit()

    asyncio.run(engine.continue_mission(task_id))
    with new_session() as session:
        assert session.get(Mission, task_id).status == "succeeded"

    # 重新校验 → 应生成新版本文档（版本递增，不覆盖旧稿）
    asyncio.run(engine.run_mission(task_id, start_from="S6"))
    directory = task_dir(task_id)
    assert (directory / "document-v1.docx").exists()
    assert (directory / "document-v2.docx").exists(), "重新校验后应生成 v2"
    report = json.loads((directory / "check-report.json").read_text(encoding="utf-8"))
    assert report["document_version"] == 2
